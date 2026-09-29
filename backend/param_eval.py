# -*- coding: utf-8 -*-
"""
Day6 参数对比实验（控制变量法）
- 固定同一批 20 题（应答15 + 拒答5，含已知 bad case：本地加载/消息类型/embedding/天气/算术）
- 向量库、模型(qwen-plus)、temperature=0.2、提示词全部不变，每次只改“检索方式”
- 4 个配置：k=2 / k=3(基线) / k=5 / Rerank(粗召回10段→gte-rerank-v2精排取3段)
- 输出 RAG参数对比.xlsx：每题每组答案 + 判定，汇总表用公式统计各组正确数
运行（backend 目录、已激活 venv）：python param_eval.py
"""
import os
from http import HTTPStatus
from dotenv import load_dotenv
from langchain_community.vectorstores import Chroma
from langchain_openai import OpenAIEmbeddings, ChatOpenAI
from langchain_core.prompts import ChatPromptTemplate, MessagesPlaceholder
import dashscope
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.worksheet.datavalidation import DataValidation

load_dotenv()
API_KEY = os.getenv("DASHSCOPE_API_KEY")
dashscope.api_key = API_KEY

# ---- 复用与 app.py / eval.py 完全一致的组件（控制变量，只有检索方式变）----
embeddings = OpenAIEmbeddings(
    model="qwen3.7-text-embedding",
    api_key=API_KEY,
    base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
    check_embedding_ctx_length=False,
    chunk_size=20,
)
db = Chroma(persist_directory="chroma_db", embedding_function=embeddings,
            collection_name="knowledge_base")
llm = ChatOpenAI(
    model="qwen-plus", temperature=0.2, api_key=API_KEY,
    base_url="https://dashscope.aliyuncs.com/compatible-mode/v1",
)
CONV_PROMPT = ChatPromptTemplate.from_messages([
    ("system",
     "你是严谨的知识库问答助手，答案必须来自参考资料，"
     "资料不足时回答：抱歉，知识库中没有相关信息；"
     "回答末尾用[1][2]标注引用。"),
    MessagesPlaceholder(variable_name="history"),
    ("human", "参考资料：\n{context}\n\n问题：{question}"),
])


def format_docs(docs):
    return "\n\n".join(f"[{i+1}] {d.page_content}" for i, d in enumerate(docs))


# ---- 固定 20 题：(问题, 类型, 期望要点) ----
CASES = [
    # 应答 15（含原评测第 9 / 17 / 20 三个 bad case）
    ("LangChain中一个最简单的LLM应用由哪几部分组成？", "应答", "提示词模板、模型、输出解析器"),
    ("什么是提示词模板（PromptTemplate）？", "应答", "模板字符串+输入变量"),
    ("LangChain如何让大模型流式输出？", "应答", ".stream()逐token输出"),
    ("什么是检索增强生成（RAG）？", "应答", "先检索再生成、减少幻觉"),
    ("文档为什么要切分成小块（chunk）？", "应答", "适配上下文、提高检索精度"),
    ("本地RAG应用通常用什么加载本地文档？", "应答", "文档加载器Document Loader"),
    ("多轮对话中如何让检索器听懂带代词的追问？", "应答", "历史感知检索器改写问题"),
    ("什么是智能体（Agent）？", "应答", "LLM推理、自主调工具、循环完成"),
    ("LangChain中有哪几种聊天消息类型？", "应答", "Human/AI/System三类消息"),
    ("什么是文本嵌入（embedding）？", "应答", "文本转向量、语义近则距离近"),
    ("MMR（最大边际相关性）检索解决什么问题？", "应答", "平衡相关性与多样性"),
    ("什么是向量存储（vector store）？", "应答", "存向量、支持相似度检索如Chroma"),
    ("信息提取（information extraction）是做什么的？", "应答", "从非结构化文本抽结构化信息"),
    ("什么是过拟合？如何缓解？", "应答", "训练好泛化差；正则/dropout/加数据"),
    ("什么是注意力机制（attention）？", "应答", "动态分配权重、聚焦相关信息"),
    # 拒答 5（含原评测第 36 / 50 两个 bad case）
    ("今天郑州天气怎么样？", "拒答", "应拒答（实时天气）"),
    ("帮我算一下23乘以17再加5等于几。", "拒答", "应拒答（纯计算越界）"),
    ("红烧肉的家常做法是什么？", "拒答", "应拒答（菜谱）"),
    ("现在人民币兑美元的汇率是多少？", "拒答", "应拒答（实时汇率）"),
    ("推荐几家郑州好吃的火锅店。", "拒答", "应拒答（本地生活）"),
]


def retrieve_k(question, k):
    """纯向量检索，直接取最相似的 k 段。"""
    return db.as_retriever(search_kwargs={"k": k}).invoke(question)


def retrieve_rerank(question, coarse_k=10, top_n=3):
    """两阶段：向量粗召回 coarse_k 段，再用 gte-rerank-v2 精排取 top_n 段。"""
    candidates = db.as_retriever(search_kwargs={"k": coarse_k}).invoke(question)
    texts = [d.page_content for d in candidates]
    resp = dashscope.TextReRank.call(
        model="gte-rerank-v2",
        query=question,
        documents=texts,
        top_n=top_n,
        return_documents=False,
    )
    if resp.status_code != HTTPStatus.OK:
        raise RuntimeError(f"rerank 返回码 {resp.status_code}: {resp.message}")
    picked = [candidates[item.index] for item in resp.output.results[:top_n]]
    return picked


# 4 个配置：名称 -> 检索函数（只有这里不同）
CONFIGS = [
    ("k=2", lambda q: retrieve_k(q, 2)),
    ("k=3(基线)", lambda q: retrieve_k(q, 3)),
    ("k=5", lambda q: retrieve_k(q, 5)),
    ("Rerank(10→3)", lambda q: retrieve_rerank(q)),
]


def answer(question, docs):
    msgs = CONV_PROMPT.format_messages(
        question=question, context=format_docs(docs), history=[])
    return llm.invoke(msgs).content


def auto_judge(typ, ans):
    """规则能确定的先判；应答类是否答对要点仍留人工（返回空）。"""
    refused = "没有相关信息" in ans
    if typ == "拒答":
        return "是" if refused else "否"      # 拒答题：确实拒答=正确
    if refused:
        return "否"                            # 应答题却拒答=错误（漏召回）
    return ""                                  # 作答了，对错人工核对


def main():
    wb = Workbook()
    ws = wb.active
    ws.title = "参数对比"

    # 表头：A-D 固定信息 + 每个配置两列（答案、判定）
    headers = ["编号", "问题", "类型", "期望要点"]
    for name, _ in CONFIGS:
        headers += [f"{name}\n答案", f"{name}\n判定(是/否)"]
    ws.append(headers)

    n = len(CASES)
    for i, (q, typ, expect) in enumerate(CASES, start=1):
        row = [i, q, typ, expect]
        print(f"[{i}/{n}] {q}")
        for name, retr in CONFIGS:
            try:
                docs = retr(q)
                ans = answer(q, docs)
                judge = auto_judge(typ, ans)
            except Exception as e:
                ans, judge = f"【调用失败】{e}", "否"
            row += [ans, judge]
        ws.append(row)

    # ---- 样式 ----
    head_fill = PatternFill("solid", fgColor="16406B")
    head_font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
    center = Alignment(horizontal="center", vertical="center", wrap_text=True)
    for c in ws[1]:
        c.fill = head_fill
        c.font = head_font
        c.alignment = center
    body_font = Font(name="微软雅黑", size=10)
    for r in range(2, n + 2):
        for c in range(1, 13):
            cell = ws.cell(row=r, column=c)
            cell.font = body_font
            cell.alignment = center

    widths = {"A": 6, "B": 30, "C": 7, "D": 24,
              "E": 48, "F": 9, "G": 48, "H": 9,
              "I": 48, "J": 9, "K": 48, "L": 9}
    for col, w in widths.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = "E2"   # 冻结首行和前4列，横向对比时题号问题常驻

    # 判定列下拉（F/H/J/L）
    dv = DataValidation(type="list", formula1='"是,否"', allow_blank=True)
    ws.add_data_validation(dv)
    for col in ("F", "H", "J", "L"):
        dv.add(f"{col}2:{col}{n+1}")

    # ---- 对比汇总（公式统计；应答判定人工补全后数字自动完整）----
    sm = wb.create_sheet("对比汇总")
    sm.merge_cells("A1:E1")
    sm["A1"] = "RAG 参数对比结果（固定20题，控制变量）"
    sm["A1"].fill = head_fill
    sm["A1"].font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=13)
    sm["A1"].alignment = center
    sm.append([])
    sm.append(["配置", "应答正确(/15)", "拒答正确(/5)", "总正确(/20)", "正确率"])
    D = "'参数对比'"
    type_rng = f"{D}!$C$2:$C${n+1}"
    judge_cols = {"k=2": "F", "k=3(基线)": "H", "k=5": "J", "Rerank(10→3)": "L"}
    r = 4
    for name, _ in CONFIGS:
        jc = judge_cols[name]
        jr = f"{D}!${jc}$2:${jc}${n+1}"
        sm.append([
            name,
            f'=COUNTIFS({type_rng},"应答",{jr},"是")',
            f'=COUNTIFS({type_rng},"拒答",{jr},"是")',
            f'=COUNTIF({jr},"是")',
            f"=D{r}/20",
        ])
        r += 1
    for c in sm[3]:
        c.fill = PatternFill("solid", fgColor="5B9BD5")
        c.font = Font(name="微软雅黑", bold=True, color="FFFFFF", size=11)
        c.alignment = center
    for rr in range(4, 8):
        for cc in range(1, 6):
            cell = sm.cell(row=rr, column=cc)
            cell.font = Font(name="微软雅黑", size=11)
            cell.alignment = center
        sm.cell(row=rr, column=5).number_format = "0.0%"
    for col, w in {"A": 16, "B": 14, "C": 14, "D": 13, "E": 10}.items():
        sm.column_dimensions[col].width = w
    sm["A10"] = ("说明：拒答题与“应答却拒答”的判定已由脚本按规则填好；"
                 "其余应答题请在「参数对比」表对照期望要点人工补判（是/否，下拉），"
                 "补全后本表正确率自动更新。重点看第6/9/10题（原bad case）在哪组被修好。")
    sm.merge_cells("A10:E13")
    sm["A10"].alignment = Alignment(horizontal="left", vertical="top", wrap_text=True)
    sm["A10"].font = Font(name="微软雅黑", size=10, color="808080")
    for rr in range(10, 14):
        sm.row_dimensions[rr].height = 22

    out = "RAG参数对比.xlsx"
    wb.save(out)
    print(f"\n完成，{n} 题 × {len(CONFIGS)} 配置，已生成：{os.path.abspath(out)}")


if __name__ == "__main__":
    main()
